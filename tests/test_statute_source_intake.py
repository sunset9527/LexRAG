from pathlib import Path

import pytest
from docx import Document

from src.ingestion.statute_source_intake import (
    SOURCE_STATUS_CURRENT,
    SOURCE_STATUS_PUBLISHED_NOT_EFFECTIVE,
    SOURCE_STATUS_SUPPLEMENTAL_AMENDMENT,
    StatuteSourceIntakeError,
    inspect_stage4_source,
    scan_stage4_source_directory,
)


def _write_docx(path: Path, title: str, body: str = "第一条 为了测试版本台账。") -> None:
    document = Document()
    document.add_paragraph(title)
    document.add_paragraph(body)
    document.save(path)


def test_inspect_stage4_source_records_current_company_law_provenance(tmp_path):
    source = tmp_path / "中华人民共和国公司法_20231229.docx"
    _write_docx(source, "中华人民共和国公司法")

    candidate = inspect_stage4_source(source)

    assert candidate.law_slug == "company-law"
    assert candidate.law_name == "中华人民共和国公司法"
    assert candidate.version_id == "company-law@2023-12-29"
    assert candidate.effective_from.isoformat() == "2024-07-01"
    assert candidate.source_status == SOURCE_STATUS_CURRENT
    assert candidate.raw_sha256
    assert candidate.content_sha256
    assert candidate.raw_sha256 != candidate.content_sha256
    assert candidate.authority_status == "PENDING_REVIEW"


def test_inspect_stage4_source_keeps_not_yet_effective_trademark_separate(tmp_path):
    source = tmp_path / "中华人民共和国商标法_20260626.docx"
    _write_docx(source, "中华人民共和国商标法")

    candidate = inspect_stage4_source(source)

    assert candidate.version_id == "trademark-law@2026-06-26"
    assert candidate.effective_from.isoformat() == "2027-01-01"
    assert candidate.source_status == SOURCE_STATUS_PUBLISHED_NOT_EFFECTIVE


def test_inspect_stage4_source_rejects_a_file_with_the_wrong_legal_title(tmp_path):
    source = tmp_path / "中华人民共和国公司法_20231229.docx"
    _write_docx(source, "中华人民共和国劳动法")

    with pytest.raises(StatuteSourceIntakeError, match="正文标题"):
        inspect_stage4_source(source)


def test_inspect_stage4_source_keeps_criminal_law_amendment_as_supplement_only(tmp_path):
    source = tmp_path / "中华人民共和国刑法修正案（十二）_20231229.docx"
    _write_docx(source, "中华人民共和国刑法修正案（十二）")

    candidate = inspect_stage4_source(source)

    assert candidate.law_slug == "criminal-law"
    assert candidate.version_id == "criminal-amendment-12@2023-12-29"
    assert candidate.document_type == "AMENDMENT"
    assert candidate.source_status == SOURCE_STATUS_SUPPLEMENTAL_AMENDMENT


def test_inspect_stage4_source_registers_constitution_and_current_public_security_law(tmp_path):
    constitution = tmp_path / "中华人民共和国宪法（2018年修正文本）_20180311.docx"
    public_security = tmp_path / "中华人民共和国治安管理处罚法_20250627.docx"
    _write_docx(constitution, "中华人民共和国宪法")
    _write_docx(public_security, "中华人民共和国治安管理处罚法")

    constitution_candidate = inspect_stage4_source(constitution)
    public_security_candidate = inspect_stage4_source(public_security)

    assert constitution_candidate.version_id == "constitution@2018-03-11"
    assert constitution_candidate.source_status == SOURCE_STATUS_CURRENT
    assert public_security_candidate.version_id == "public-security-administration-punishment-law@2025-06-27"
    assert public_security_candidate.source_status == SOURCE_STATUS_CURRENT


def test_scan_stage4_source_directory_returns_registered_candidates_and_skips_other_files(tmp_path):
    _write_docx(tmp_path / "中华人民共和国公司法_20231229.docx", "中华人民共和国公司法")
    _write_docx(tmp_path / "中华人民共和国宪法（2018年修正文本）_20180311.docx", "中华人民共和国宪法")
    _write_docx(tmp_path / "中华人民共和国治安管理处罚法_20250627.docx", "中华人民共和国治安管理处罚法")

    scan = scan_stage4_source_directory(tmp_path)

    assert [candidate.version_id for candidate in scan.candidates] == [
        "company-law@2023-12-29",
        "constitution@2018-03-11",
        "public-security-administration-punishment-law@2025-06-27",
    ]
    assert scan.skipped_files == ()
