"""通用文档解析器测试。"""

import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

from src.ingestion.document_parser import parse_document


def test_parse_docx_keeps_paragraphs_and_table_cells(tmp_path):
    """Word 解析必须保留正文段落和表格单元格，避免证据材料信息缺失。"""
    source = tmp_path / "劳动仲裁证据.docx"
    document = Document()
    document.add_paragraph("劳动者于2025年1月入职。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "岗位"
    table.cell(0, 1).text = "研发工程师"
    document.save(source)

    segments = parse_document(source)

    assert [segment["text"] for segment in segments] == [
        "劳动者于2025年1月入职。",
        "岗位 | 研发工程师",
    ]
    assert [segment["paragraph"] for segment in segments] == [1, 2]


def test_parse_pdf_preserves_page_number_for_source_citation():
    """PDF 解析结果必须保留页码，供检索结果定位原始来源。"""
    source = ROOT / "tests" / "fixtures" / "sample_text_law.pdf"

    segments = parse_document(source)

    assert len(segments) == 1
    assert segments[0]["page"] == 1
    assert segments[0]["text"].strip()
