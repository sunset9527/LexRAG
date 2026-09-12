"""PDF 与 Word 文档统一解析为可分块的文本段。"""

from pathlib import Path
from typing import Any


def parse_document(path: str | Path) -> list[dict[str, Any]]:
    """按文件类型提取文本，并保留页码或段落位置。"""
    document_path = Path(path)
    if document_path.suffix.lower() == ".docx":
        return _parse_docx(document_path)
    if document_path.suffix.lower() == ".pdf":
        return _parse_pdf(document_path)
    raise ValueError("仅支持 PDF 或 DOCX 文件")


def _parse_docx(path: Path) -> list[dict[str, Any]]:
    from docx import Document

    document = Document(path)
    segments: list[dict[str, Any]] = []
    paragraph_number = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraph_number += 1
            segments.append({"text": text, "paragraph": paragraph_number})
    for table in document.tables:
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells).strip()
            if text:
                paragraph_number += 1
                segments.append({"text": text, "paragraph": paragraph_number})
    return segments


def _parse_pdf(path: Path) -> list[dict[str, Any]]:
    from src.ingestion.pdf_parser import parse_pdf

    return [
        {"text": page["text"].strip(), "page": page["page"], "source": page["source"]}
        for page in parse_pdf(path)
        if page["text"].strip()
    ]
